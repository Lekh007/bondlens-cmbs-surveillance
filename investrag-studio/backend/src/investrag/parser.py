from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import mimetypes
import re
import zipfile
from collections.abc import Iterable
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from defusedxml import ElementTree

from .domain import CanonicalElement, ParsedDocument


class UnsupportedFormatError(ValueError):
    pass


def parser_capabilities() -> list[dict[str, object]]:
    """Expose optional parser availability to the UI and demo script."""

    return [
        {"name": "docling", "installed": bool(importlib.util.find_spec("docling")), "role": "primary layout parser"},
        {"name": "mineru", "installed": bool(importlib.util.find_spec("mineru")), "role": "difficult-scan escalation"},
        {"name": "pymupdf", "installed": bool(importlib.util.find_spec("fitz")), "role": "native PDF validation"},
        {"name": "openpyxl", "installed": bool(importlib.util.find_spec("openpyxl")), "role": "workbook structure"},
        {"name": "beautifulsoup4", "installed": bool(importlib.util.find_spec("bs4")), "role": "HTML structure"},
    ]


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self.skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self.skip_depth and data.strip():
            self.parts.append(data.strip())


def detect_format(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    by_suffix = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".csv": "text/csv",
        ".html": "text/html",
        ".htm": "text/html",
        ".json": "application/json",
        ".md": "text/markdown",
        ".txt": "text/plain",
        ".eml": "message/rfc822",
        ".msg": "application/vnd.ms-outlook",
        ".zip": "application/zip",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
    }
    media_type = by_suffix.get(suffix) or mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return suffix.lstrip(".") or "unknown", media_type


def _element_id(source_name: str, index: int, text: str) -> str:
    digest = hashlib.sha1(f"{source_name}:{index}:{text}".encode()).hexdigest()[:12]
    return f"el_{digest}"


def _elements(source_name: str, rows: Iterable[tuple[str, str, dict[str, Any]]]) -> list[CanonicalElement]:
    result: list[CanonicalElement] = []
    for index, (element_type, text, metadata) in enumerate(rows):
        clean = re.sub(r"\s+", " ", text).strip()
        if not clean:
            continue
        result.append(
            CanonicalElement(
                element_id=_element_id(source_name, index, clean),
                element_type=element_type,  # type: ignore[arg-type]
                text=clean,
                order=index,
                page=metadata.pop("page", None),
                slide=metadata.pop("slide", None),
                sheet=metadata.pop("sheet", None),
                cell_range=metadata.pop("cell_range", None),
                json_path=metadata.pop("json_path", None),
                bbox=metadata.pop("bbox", None),
                confidence=metadata.pop("confidence", 1.0),
                warnings=metadata.pop("warnings", []),
                metadata=metadata,
            )
        )
    return result


def _parse_text(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    rows: list[tuple[str, str, dict[str, Any]]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            rows.append(("heading", stripped.lstrip("# "), {}))
        elif stripped.startswith(('-', '*')):
            rows.append(("list", stripped[1:].strip(), {}))
        else:
            rows.append(("paragraph", stripped, {}))
    return ParsedDocument(
        source_name=source_name,
        media_type=media_type,
        parser="native-text",
        parser_version="1.0",
        elements=_elements(source_name, rows),
    )


def _parse_html(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw, "html.parser")
        rows: list[tuple[str, str, dict[str, Any]]] = []
        for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th"]):
            kind = "heading" if node.name.startswith("h") else "list" if node.name == "li" else "paragraph"
            rows.append((kind, node.get_text(" ", strip=True), {}))
        parser_name = "beautifulsoup4"
    except ImportError:
        fallback = _TextHTMLParser()
        fallback.feed(raw)
        rows = [("paragraph", text, {}) for text in fallback.parts]
        parser_name = "stdlib-html-parser"
    return ParsedDocument(
        source_name=source_name,
        media_type=media_type,
        parser=parser_name,
        parser_version="1.0",
        elements=_elements(source_name, rows),
    )


def _parse_json(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows: list[tuple[str, str, dict[str, Any]]] = []

    def visit(value: object, location: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                visit(child, f"{location}.{key}")
        elif isinstance(value, list):
            for index, child in enumerate(value):
                visit(child, f"{location}[{index}]")
        else:
            rows.append(("metadata", f"{location}: {value}", {"json_path": location}))

    visit(payload, "$")
    return ParsedDocument(source_name=source_name, media_type=media_type, parser="native-json", parser_version="1.0", elements=_elements(source_name, rows))


def _parse_csv(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    rows: list[tuple[str, str, dict[str, Any]]] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, None)
        if header:
            rows.append(("table", " | ".join(header), {"cell_range": "header"}))
        for row_number, row in enumerate(reader, start=2):
            rows.append(("table_row", " | ".join(row), {"cell_range": f"row:{row_number}"}))
    return ParsedDocument(source_name=source_name, media_type=media_type, parser="native-csv", parser_version="1.0", elements=_elements(source_name, rows))


def _parse_xlsx(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedFormatError("openpyxl is required for XLSX ingestion") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    rows: list[tuple[str, str, dict[str, Any]]] = []
    for sheet in workbook.worksheets:
        rows.append(("heading", f"Sheet: {sheet.title}", {"sheet": sheet.title}))
        for row in sheet.iter_rows():
            values = ["" if cell.value is None else str(cell.value) for cell in row]
            if any(values):
                first = row[0].row
                last = row[-1].column_letter
                rows.append(("spreadsheet_range", " | ".join(values), {"sheet": sheet.title, "cell_range": f"A{first}:{last}{first}"}))
    return ParsedDocument(source_name=source_name, media_type=media_type, parser="openpyxl", parser_version="3", elements=_elements(source_name, rows))


def _parse_pdf(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedFormatError("PyMuPDF is required for PDF ingestion") from exc
    document = fitz.open(path)
    rows: list[tuple[str, str, dict[str, Any]]] = []
    warnings: list[str] = []
    for page_index, page in enumerate(document, start=1):
        blocks = page.get_text("blocks")
        if not blocks:
            warnings.append(f"page {page_index} has no native text; OCR escalation recommended")
            continue
        for block in blocks:
            text = block[4].strip()
            if text:
                rows.append(("paragraph", text, {"page": page_index, "bbox": tuple(block[:4])}))
    quality = 1.0 if rows else 0.0
    if warnings:
        quality = max(0.0, quality - min(0.4, 0.05 * len(warnings)))
    return ParsedDocument(source_name=source_name, media_type=media_type, parser="pymupdf-native", parser_version="1.26", elements=_elements(source_name, rows), warnings=warnings, quality_score=quality)


def _parse_office_xml(path: Path, source_name: str, media_type: str, kind: str) -> ParsedDocument:
    rows: list[tuple[str, str, dict[str, Any]]] = []
    warnings: list[str] = []
    with zipfile.ZipFile(path) as archive:
        if kind == "docx":
            names = [name for name in archive.namelist() if name == "word/document.xml"]
            if not names:
                raise UnsupportedFormatError("DOCX has no word/document.xml")
            root = ElementTree.fromstring(archive.read(names[0]))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for paragraph in root.findall(".//w:p", ns):
                text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
                rows.append(("paragraph", text, {}))
            parser_name = "docx-xml-fallback"
        else:
            names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for slide_index, name in enumerate(names, start=1):
                root = ElementTree.fromstring(archive.read(name))
                text = " ".join(node.text or "" for node in root.findall(".//a:t", ns))
                rows.append(("paragraph", text, {"slide": slide_index}))
            parser_name = "pptx-xml-fallback"
            if not names:
                warnings.append("presentation contains no slide XML")
    return ParsedDocument(source_name=source_name, media_type=media_type, parser=parser_name, parser_version="1.0", elements=_elements(source_name, rows), warnings=warnings, quality_score=0.9 if rows else 0.0)


def _parse_email(path: Path, source_name: str, media_type: str) -> ParsedDocument:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    rows: list[tuple[str, str, dict[str, Any]]] = []
    for field in ("Subject", "From", "To", "Date"):
        value = message.get(field)
        if value:
            rows.append(("metadata", f"{field}: {value}", {}))
    body = message.get_body(preferencelist=("plain", "html"))
    if body:
        rows.append(("email", body.get_content(), {}))
    attachments = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
    warnings = [f"attachment requires recursive ingestion: {name}" for name in attachments]
    return ParsedDocument(source_name=source_name, media_type=media_type, parser="email-parser", parser_version="1.0", elements=_elements(source_name, rows), warnings=warnings, quality_score=0.85 if body else 0.5)


def parse_document(path: Path, source_name: str | None = None) -> ParsedDocument:
    """Parse a document into canonical elements with location provenance."""

    source_name = source_name or path.name
    kind, media_type = detect_format(path)
    if kind in {"txt", "md"}:
        return _parse_text(path, source_name, media_type)
    if kind in {"html", "htm"}:
        return _parse_html(path, source_name, media_type)
    if kind == "json":
        return _parse_json(path, source_name, media_type)
    if kind == "csv":
        return _parse_csv(path, source_name, media_type)
    if kind == "xlsx":
        return _parse_xlsx(path, source_name, media_type)
    if kind == "pdf":
        return _parse_pdf(path, source_name, media_type)
    if kind in {"docx", "pptx"}:
        return _parse_office_xml(path, source_name, media_type, kind)
    if kind in {"eml", "msg"}:
        return _parse_email(path, source_name, media_type)
    if kind in {"png", "jpg", "jpeg", "tif", "tiff"}:
        return ParsedDocument(source_name=source_name, media_type=media_type, parser="image-placeholder", parser_version="1.0", elements=[], warnings=["OCR is required before this image can be trusted"], quality_score=0.0)
    if kind == "zip":
        raise UnsupportedFormatError("ZIP files must be safely extracted by the ingestion service")
    raise UnsupportedFormatError(f"unsupported document format: {media_type}")
